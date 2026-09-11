from __future__ import annotations

import json
import shutil
import subprocess
import sys
import threading
import csv
import io
import re
from html import escape as html_escape
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from pathlib import Path
from typing import Any

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .adapters import AdapterError, NoticeData, make_adapter
from .config import normalize_data_dir, settings
from .db import beijing_time, get_db, init_db, json_load, log_event, now_iso, recover_incomplete_runs, select_site_account
from .parsers import file_mime, parse_document, sha256_file, parser_capabilities
from .service import create_attachment, frontend_notice, ingest_notice_data, ingest_url, permanently_delete_notices, rebuild_keyword_group_analysis, refresh_notice_analysis, reparse_notice, run_crawl, try_create_run
from .scheduler import scheduler
from .storage import absolute_from_relative, attachment_directory, relative_to_data, safe_name
from .stabilization_migration import migrate_stabilization_schema
from . import reparse_tasks
from .maintenance import activity
from .migration import adopt_storage, migrate_storage
from .update_checker import update_monitor
from .manual_verification import ManualVerificationError, browser_session_port, close_verification, complete_verification, open_verification


APP_VERSION = "1.2.4"
app = FastAPI(title="猎标 V1 API", version=APP_VERSION, docs_url="/api/docs", redoc_url=None)
PARSE_ISSUE_SQL = """(n.ingest_status='failed' OR EXISTS (
    SELECT 1 FROM attachments ia WHERE ia.notice_id=n.id
    AND (ia.parse_status IN ('failed','unsupported','ocr_pending','pending')
         OR ia.status IN ('failed','needs_tool','not_downloaded','downloading'))
))"""
BASE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))


@app.middleware("http")
async def prevent_stale_frontend(request, call_next):
    try:
        if request.method == 'PATCH' and request.url.path == '/api/settings/storage':
            response = await call_next(request)
        else:
            with activity():
                response = await call_next(request)
    except RuntimeError as exc:
        return JSONResponse(status_code=503, content={'detail': str(exc)})
    if request.url.path in ("", "/") or request.url.path.endswith((".js", ".css")):
        response.headers["Cache-Control"] = "no-store, max-age=0"
    return response


class MarkRequest(BaseModel):
    mark: str = Field(pattern="^(pending|relevant|irrelevant|supplement|focus|processed)$")


class BatchActionRequest(BaseModel):
    ids: list[int] = Field(min_length=1, max_length=500)
    mark: str | None = Field(default=None, pattern="^(pending|relevant|irrelevant|supplement|focus|processed)$")


class FieldCorrectionRequest(BaseModel):
    value: str = Field(max_length=2000)


class KeywordGroupRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    category: str | None = Field(default=None, max_length=100)
    enabled: bool = True
    include_any: list[str] = Field(default_factory=list)
    include_all: list[str] = Field(default_factory=list)
    phrases: list[str] = Field(default_factory=list)
    exclude: list[str] = Field(default_factory=list)
    synonyms: dict[str, list[str]] = Field(default_factory=dict)
    scopes: list[str] = Field(default_factory=lambda: ["title", "body", "attachment_name", "attachment_body"])
    priority: int = Field(default=0, ge=-1000, le=1000)


class KeyFileRequest(BaseModel):
    is_key_file: bool


class AccountRequest(BaseModel):
    alias: str
    username: str | None = None
    login_mode: str = "manual_session"
    credential_ref: str | None = None
    enabled: bool = True


class UrlImportRequest(BaseModel):
    urls: list[str] = Field(min_length=1, max_length=100)
    site_code: str | None = None


class StorageRequest(BaseModel):
    root: str | None = None
    max_attachment_mb: int | None = Field(default=None, ge=1, le=4096)
    max_archive_mb: int | None = Field(default=None, ge=1, le=4096)
    overwrite: bool = False


class JobRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    site_id: int
    account_id: int | None = None
    keyword_group_id: int | None = None
    schedule_text: str = "每 30 分钟"
    categories: list[str] = Field(default_factory=lambda: ["招标公告"])
    timezone: str = "Asia/Shanghai"
    lookback_days: int = Field(default=1, ge=0, le=3650)
    max_pages: int = Field(default=5, ge=1, le=100)
    max_notices: int = Field(default=100, ge=1, le=10000)
    concurrency: int = Field(default=1, ge=1, le=8)
    interval_ms: int = Field(default=1500, ge=200, le=60000)
    retry_max_attempts: int = Field(default=3, ge=0, le=10)
    download_attachments: bool = True
    ocr_enabled: bool = False
    enabled: bool = True


class EnabledRequest(BaseModel):
    enabled: bool


class RunJobRequest(BaseModel):
    lookback_days: int | None = Field(default=None, ge=0, le=3650)
    max_pages: int | None = Field(default=None, ge=1, le=100)
    max_notices: int | None = Field(default=None, ge=1, le=10000)
    download_attachments: bool | None = None


class RunAllJobsRequest(BaseModel):
    lookback_days: int | None = Field(default=None, ge=0, le=3650)
    max_pages: int | None = Field(default=None, ge=1, le=100)
    max_notices: int | None = Field(default=None, ge=1, le=10000)


def validate_id(value: int, label: str = "记录") -> None:
    if value <= 0:
        raise HTTPException(status_code=400, detail=f"{label} ID 无效")


def validate_job_request(payload: JobRequest) -> None:
    payload.categories = list(dict.fromkeys(item.strip() for item in payload.categories if item.strip()))
    if not payload.categories:
        raise HTTPException(status_code=400, detail="至少选择一种公告类型")
    if payload.concurrency != 1:
        raise HTTPException(status_code=400, detail="当前版本仅支持单任务串行采集，并发数必须为 1")
    if payload.ocr_enabled:
        from .ocr import available
        if not available():
            raise HTTPException(status_code=400, detail="当前程序未包含 CPU OCR 组件，无法启用自动 OCR")


def localize(item: dict[str, Any], fields: tuple[str, ...]) -> dict[str, Any]:
    for field in fields:
        if field in item:
            item[field] = beijing_time(item[field])
    return item


@app.on_event("startup")
def startup() -> None:
    settings.load_persisted_data_dir()
    settings.ensure_dirs()
    init_db()
    migrate_stabilization_schema()
    recover_incomplete_runs()
    reparse_tasks.initialize(recover=True)
    scheduler.start()
    update_monitor.start(APP_VERSION, settings.update_url, settings.update_enabled)


@app.on_event("shutdown")
def shutdown() -> None:
    update_monitor.stop()
    scheduler.stop()


@app.get("/api/health")
def health() -> dict[str, Any]:
    with get_db() as connection:
        connection.execute("SELECT 1").fetchone()
    return {"ok": True, "service": "lieBiao", "version": APP_VERSION, "timezone": "Asia/Shanghai", "timezone_label": "北京时间（UTC+08:00）", "server_time": beijing_time(now_iso()), "storage_root": str(settings.data_dir), "database": str(settings.db_path)}


@app.get("/api/settings/storage")
def get_storage_settings() -> dict[str, Any]:
    with get_db() as connection:
        rows = connection.execute("SELECT key,value FROM app_settings WHERE key LIKE 'storage.%'").fetchall()
    values = {row["key"].removeprefix("storage."): row["value"] for row in rows}
    return {"root": str(settings.data_dir), "raw": str(settings.raw_dir), "extracted": str(settings.extracted_dir), "preview": str(settings.preview_dir), "temp": str(settings.temp_dir), "logs": str(settings.log_dir), "max_attachment_mb": int(values.get("max_attachment_mb", settings.max_attachment_mb)), "max_archive_mb": int(values.get("max_archive_mb", settings.max_archive_mb))}


@app.get('/api/settings/parsers')
def get_parser_capabilities():
    return parser_capabilities()


@app.get("/api/update/check")
def get_update_status() -> dict[str, Any]:
    return update_monitor.status()


def _copy_storage_tree(source: Path, target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for child in source.iterdir():
        destination = target / child.name
        if child.is_dir():
            shutil.copytree(child, destination, dirs_exist_ok=True)
        else:
            shutil.copy2(child, destination)


@app.patch("/api/settings/storage")
def update_storage_settings(payload: StorageRequest) -> dict[str, Any]:
    timestamp = now_iso()
    requested_root = payload.root.strip() if payload.root else ""
    if requested_root:
        try:
            root = normalize_data_dir(requested_root)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
        old_root = settings.data_dir.resolve()
        if root != old_root:
            if root == Path(root.anchor):
                raise HTTPException(status_code=400, detail="文件目录不能直接使用磁盘根目录")
            try:
                root.relative_to(old_root)
                nested = True
            except ValueError:
                try:
                    old_root.relative_to(root)
                    nested = True
                except ValueError:
                    nested = False
            if nested:
                raise HTTPException(status_code=400, detail="新目录不能位于当前目录内部或作为当前目录的上级目录")
            if not settings.db_path.exists():
                raise HTTPException(status_code=500, detail="当前数据库不存在，无法迁移目录")
            target_has_data = root.exists() and any(root.iterdir())
            if target_has_data:
                if not (root / "scout.db").is_file():
                    raise HTTPException(status_code=400, detail="目标目录非空且不是已有猎标数据目录")
                if not payload.overwrite:
                    raise HTTPException(status_code=409, detail="检测到已有猎标数据库。切换后不会覆盖其内容，是否确认使用该目录？")
                scheduler.stop()
                try:
                    moved_from = adopt_storage(root)
                    init_db()
                    migrate_stabilization_schema()
                    reparse_tasks.initialize(recover=True)
                except ValueError as exc:
                    raise HTTPException(status_code=400, detail=str(exc)) from exc
                except RuntimeError as exc:
                    raise HTTPException(status_code=409, detail=str(exc)) from exc
                finally:
                    scheduler.start()
            backup_db = None
            if not (root / "scout.db").is_file():
                try:
                    moved_from = migrate_storage(root)
                except ValueError as exc:
                    raise HTTPException(status_code=400, detail=str(exc)) from exc
                except RuntimeError as exc:
                    raise HTTPException(status_code=409, detail=str(exc)) from exc
        else:
            settings.persist_data_dir()
    settings.ensure_dirs()
    values = {
        "storage.root": str(settings.data_dir),
        "storage.raw": str(settings.raw_dir),
        "storage.extracted": str(settings.extracted_dir),
        "storage.preview": str(settings.preview_dir),
        "storage.temp": str(settings.temp_dir),
        "storage.logs": str(settings.log_dir),
    }
    if payload.max_attachment_mb is not None:
        values["storage.max_attachment_mb"] = str(payload.max_attachment_mb)
    if payload.max_archive_mb is not None:
        values["storage.max_archive_mb"] = str(payload.max_archive_mb)
    with get_db() as connection:
        for key, value in values.items():
            connection.execute("INSERT INTO app_settings(key,value,updated_at) VALUES(?,?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value,updated_at=excluded.updated_at", (key, value, timestamp))
        log_event(connection, "settings.storage", f"更新文件目录：{settings.data_dir}")
    settings.persist_data_dir()
    response = get_storage_settings()
    if requested_root and "moved_from" in locals():
        response["moved_from"] = moved_from
        response["database_backup"] = str(backup_db) if backup_db else None
    return response


@app.get("/api/sites")
def list_sites() -> dict[str, Any]:
    with get_db() as connection:
        sites = [localize(dict(row), ("last_checked_at", "created_at")) for row in connection.execute("SELECT * FROM sites WHERE enabled=1 ORDER BY id").fetchall()]
        for site in sites:
            accounts = connection.execute("SELECT id,alias,username,login_mode,session_status,last_login_at,expires_at,status_reason,enabled FROM site_accounts WHERE site_id=? ORDER BY id", (site["id"],)).fetchall()
            site["accounts"] = [localize(dict(account), ("last_login_at", "expires_at", "created_at")) for account in accounts]
    return {"items": sites}


@app.post("/api/sites/{site_id}/health-check")
def site_health_check(site_id: int) -> dict[str, Any]:
    validate_id(site_id, "平台")
    with get_db() as connection:
        site = connection.execute("SELECT * FROM sites WHERE id=?", (site_id,)).fetchone()
        account = select_site_account(connection, site_id)
    if not site:
        raise HTTPException(status_code=404, detail="平台不存在")
    adapter = make_adapter(site["code"], site["base_url"], session_cookie=account["credential_ref"] if account else None)
    try:
        result = adapter.health_check()
    finally:
        adapter.close()
    with get_db() as connection:
        connection.execute("UPDATE sites SET health_status=?,health_message=?,last_checked_at=? WHERE id=?", ("healthy" if result["ok"] else "unhealthy", result["message"], now_iso(), site_id))
        log_event(connection, "site.health", f"{site['name']}：{result['message']}", "INFO" if result["ok"] else "WARNING")
    return result


@app.post("/api/sites/{site_id}/manual-verification/open")
def open_site_verification(site_id: int) -> dict[str, Any]:
    validate_id(site_id, "平台")
    with get_db() as connection:
        site = connection.execute("SELECT id,name,code,base_url FROM sites WHERE id=?", (site_id,)).fetchone()
    if not site:
        raise HTTPException(status_code=404, detail="平台不存在")
    try:
        verification_url = site["base_url"]
        if site["code"] == "chng":
            verification_url = "https://ec.chng.com.cn/channel/home/#/purchase?top=0"
        elif site["code"] == "yfb":
            verification_url = "https://qiye.qianlima.com/new_qd_yfbsite/#/infoCenter/search"
        elif site["code"] == "espic":
            verification_url = "https://ebid.espic.com.cn/newgdtcms//category/bulletinListNew.html?dates=300&categoryId=2&tenderMethod=01&tabName=%E6%8B%9B%E6%A0%87%E4%BF%A1%E6%81%AF&page=1"
        elif site["code"] == "chdtp":
            verification_url = "https://www.chdtp.com/pages/wzglS/cgxx/caigou.jsp?cgtype=4"
        result = open_verification(site_id, verification_url, settings.data_dir / "browser_sessions")
        # Keep the local debug port in the account record immediately. This is
        # what lets the API reconnect to an Edge window after an app restart,
        # even when the user has not clicked “验证完成” yet.
        port = result.get("port") if isinstance(result, dict) else None
        if isinstance(port, int) and 1 <= port <= 65535:
            with get_db() as connection:
                timestamp = now_iso()
                credential = f"__scout_browser_port={port}"
                account = connection.execute(
                    "SELECT id FROM site_accounts WHERE site_id=? AND alias='人工验证会话' ORDER BY id LIMIT 1",
                    (site_id,),
                ).fetchone()
                if account:
                    connection.execute(
                        "UPDATE site_accounts SET credential_ref=?,session_status='needs_manual',status_reason=?,enabled=0 WHERE id=?",
                        (credential, "验证窗口已打开，请完成验证", account["id"]),
                    )
                else:
                    connection.execute(
                        "INSERT INTO site_accounts(site_id,alias,login_mode,credential_ref,session_status,status_reason,enabled,created_at) VALUES(?,?,'manual_session',?,'needs_manual',?,0,?)",
                        (site_id, "人工验证会话", credential, "验证窗口已打开，请完成验证", timestamp),
                    )
                connection.execute(
                    "UPDATE sites SET health_status='unknown',health_message=?,last_checked_at=? WHERE id=?",
                    ("验证窗口已打开，请完成验证", timestamp, site_id),
                )
        return result
    except ManualVerificationError as exc:
        with get_db() as connection:
            timestamp = now_iso()
            message = str(exc)
            connection.execute(
                "UPDATE sites SET health_status='unhealthy',health_message=?,last_checked_at=? WHERE id=?",
                (message, timestamp, site_id),
            )
            connection.execute(
                "UPDATE site_accounts SET session_status='needs_manual',status_reason=? WHERE site_id=? AND alias='人工验证会话'",
                (message, site_id),
            )
            log_event(connection, "site.manual_verify", f"{site['name']}：{message}", "WARNING")
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/sites/{site_id}/manual-verification/complete")
def complete_site_verification(site_id: int) -> dict[str, Any]:
    validate_id(site_id, "平台")
    with get_db() as connection:
        site = connection.execute("SELECT * FROM sites WHERE id=?", (site_id,)).fetchone()
    if not site:
        raise HTTPException(status_code=404, detail="平台不存在")
    completed = False
    keep_browser = site["code"] in ("chng", "cdt", "chdtp")
    try:
        try:
            browser_port = None
            with get_db() as connection:
                account = connection.execute(
                    "SELECT credential_ref FROM site_accounts WHERE site_id=? AND alias='人工验证会话' ORDER BY id LIMIT 1",
                    (site_id,),
                ).fetchone()
                if account:
                    browser_port = browser_session_port(account["credential_ref"])
            # YFB stores its effective login token in browser storage. A valid
            # member session therefore may have no cookie at all; the browser
            # identity and debug port are the actual credential.
            if site["code"] == "yfb" and browser_port is not None:
                cookie = (
                    f"__scout_browser_port={browser_port}; "
                    f"__scout_browser_session={site_id}"
                )
            elif browser_port is None:
                cookie = complete_verification(site_id)
            else:
                cookie = complete_verification(site_id, port=browser_port)
        except ManualVerificationError as session_error:
            # Public platforms may never set a login cookie. In that case the
            # useful result is a live collection check, not a fake account.
            adapter = make_adapter(site["code"], site["base_url"])
            try:
                public_result = adapter.health_check()
            finally:
                adapter.close()
            if public_result["ok"]:
                with get_db() as connection:
                    timestamp = now_iso()
                    message = "公开采集正常，无需人工验证"
                    connection.execute(
                        "UPDATE sites SET health_status='healthy',health_message=?,last_checked_at=? WHERE id=?",
                        (message, timestamp, site_id),
                    )
                    connection.execute(
                        "UPDATE site_accounts SET session_status='public',status_reason=?,enabled=0 WHERE site_id=? AND alias='人工验证会话'",
                        (message, site_id),
                    )
                    log_event(connection, "site.manual_verify", f"{site['name']}：{message}")
                completed = True
                return {"ok": True, "mode": "public", "message": message}
            raise ManualVerificationError(f"{session_error}；公开采集检查也未通过：{public_result['message']}") from session_error
        if site["code"] in ("chng", "cdt", "yfb", "chdtp") and "__scout_browser_session=" not in cookie:
            cookie = f"{cookie}; __scout_browser_session={site_id}"
        adapter = make_adapter(site["code"], site["base_url"], session_cookie=cookie)
        try:
            result = adapter.health_check()
        finally:
            adapter.close()
        if not result["ok"]:
            raise ManualVerificationError(f"验证后仍无法采集：{result['message']}")
        if site["code"] == "yfb":
            keep_browser = result.get("mode") == "member"
        if site["code"] == "yfb" and not keep_browser:
            with get_db() as connection:
                timestamp = now_iso()
                message = "未检测到有效会员权限，已自动切换为公开采集（正文可能不完整）"
                cursor = connection.execute(
                    "UPDATE site_accounts SET credential_ref=NULL,session_status='public',last_login_at=?,status_reason=?,enabled=0 WHERE site_id=? AND alias='人工验证会话'",
                    (timestamp, message, site_id),
                )
                if cursor.rowcount == 0:
                    connection.execute(
                        "INSERT INTO site_accounts(site_id,alias,login_mode,session_status,last_login_at,status_reason,enabled,created_at) VALUES(?,?,'manual_session','public',?,?,0,?)",
                        (site_id, "人工验证会话", timestamp, message, timestamp),
                    )
                connection.execute("UPDATE crawl_jobs SET account_id=NULL WHERE site_id=?", (site_id,))
                connection.execute(
                    "UPDATE sites SET health_status='healthy',health_message=?,last_checked_at=? WHERE id=?",
                    (message, timestamp, site_id),
                )
                log_event(connection, "site.manual_verify", f"{site['name']}：{message}")
            completed = True
            return {"ok": True, "mode": "public", "message": message}
        with get_db() as connection:
            account = connection.execute(
                "SELECT id FROM site_accounts WHERE site_id=? AND alias='人工验证会话' ORDER BY id LIMIT 1", (site_id,)
            ).fetchone()
            timestamp = now_iso()
            if account:
                account_id = int(account["id"])
                connection.execute(
                    "UPDATE site_accounts SET credential_ref=?,session_status='verified',last_login_at=?,status_reason='人工验证成功',enabled=1 WHERE id=?",
                    (cookie, timestamp, account_id),
                )
            else:
                cursor = connection.execute(
                    "INSERT INTO site_accounts(site_id,alias,login_mode,credential_ref,session_status,last_login_at,status_reason,enabled,created_at) VALUES(?,?,'manual_session',?,'verified',?,'人工验证成功',1,?)",
                    (site_id, "人工验证会话", cookie, timestamp, timestamp),
                )
                account_id = int(cursor.lastrowid)
            connection.execute("UPDATE crawl_jobs SET account_id=? WHERE site_id=?", (account_id, site_id))
            connection.execute(
                "UPDATE sites SET health_status='healthy',health_message=?,last_checked_at=? WHERE id=?",
                (result.get("message") or "人工验证成功", timestamp, site_id),
            )
            log_event(connection, "site.manual_verify", f"{site['name']}：人工验证成功，会话已绑定采集任务")
        completed = True
        message = "人工验证成功，已绑定该平台采集任务，可立即采集"
        if keep_browser:
            message += "；该平台定时采集依赖专用窗口，请保持窗口打开"
        return {"ok": True, "mode": result.get("mode"), "message": message}
    except ManualVerificationError as exc:
        with get_db() as connection:
            timestamp = now_iso()
            message = str(exc)
            connection.execute(
                "UPDATE sites SET health_status='unhealthy',health_message=?,last_checked_at=? WHERE id=?",
                (message, timestamp, site_id),
            )
            connection.execute(
                "UPDATE site_accounts SET session_status='needs_manual',status_reason=? WHERE site_id=? AND alias='人工验证会话'",
                (message, site_id),
            )
            log_event(connection, "site.manual_verify", f"{site['name']}：{message}", "WARNING")
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    finally:
        if completed and not keep_browser:
            close_verification(site_id)


@app.post("/api/sites/{site_id}/accounts")
def create_account(site_id: int, payload: AccountRequest) -> dict[str, Any]:
    validate_id(site_id, "平台")
    with get_db() as connection:
        if not connection.execute("SELECT id FROM sites WHERE id=?", (site_id,)).fetchone():
            raise HTTPException(status_code=404, detail="平台不存在")
        cursor = connection.execute("INSERT INTO site_accounts(site_id,alias,username,login_mode,credential_ref,enabled,created_at) VALUES(?,?,?,?,?,?,?)", (site_id, payload.alias, payload.username, payload.login_mode, payload.credential_ref, int(payload.enabled), now_iso()))
        account_id = int(cursor.lastrowid)
        log_event(connection, "account.create", f"新增平台账号：{payload.alias}")
    return {"id": account_id, "message": "平台账号已保存；凭据仅作为引用保存"}


@app.post("/api/site-accounts/{account_id}/verify")
def verify_account(account_id: int) -> dict[str, Any]:
    validate_id(account_id, "账号")
    with get_db() as connection:
        account = connection.execute("SELECT * FROM site_accounts WHERE id=?", (account_id,)).fetchone()
        if not account:
            raise HTTPException(status_code=404, detail="账号不存在")
        if account["login_mode"] in ("captcha", "sms", "ca", "manual_session"):
            status = "needs_manual"
            reason = "需要人工完成登录/验证码后，再导入授权会话"
        else:
            status = "verified"
            reason = "公开访问模式已验证"
        connection.execute("UPDATE site_accounts SET session_status=?,last_login_at=?,status_reason=? WHERE id=?", (status, now_iso(), reason, account_id))
        log_event(connection, "account.verify", f"账号验证结果：{status}")
    return {"ok": status == "verified", "status": status, "message": reason}


@app.get("/api/keyword-groups")
def list_keyword_groups() -> dict[str, Any]:
    with get_db() as connection:
        rows = connection.execute("SELECT * FROM keyword_groups ORDER BY priority DESC,id").fetchall()
    items = []
    for row in rows:
        item = dict(row)
        for key in ("include_any_json", "include_all_json", "phrases_json", "exclude_json", "synonyms_json", "scopes_json"):
            item[key.removesuffix("_json")] = json_load(item.pop(key), [] if key != "synonyms_json" else {})
        items.append(item)
    return {"items": items}


@app.post("/api/keyword-groups")
def create_keyword_group(payload: KeywordGroupRequest) -> dict[str, Any]:
    timestamp = now_iso()
    with get_db() as connection:
        try:
            cursor = connection.execute(
                """INSERT INTO keyword_groups(name,category,enabled,include_any_json,include_all_json,phrases_json,exclude_json,synonyms_json,scopes_json,priority,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                (payload.name, payload.category, int(payload.enabled), json.dumps(payload.include_any, ensure_ascii=False),
                 json.dumps(payload.include_all, ensure_ascii=False), json.dumps(payload.phrases, ensure_ascii=False),
                 json.dumps(payload.exclude, ensure_ascii=False), json.dumps(payload.synonyms, ensure_ascii=False),
                 json.dumps(payload.scopes, ensure_ascii=False), payload.priority, timestamp, timestamp),
            )
        except Exception as exc:
            raise HTTPException(status_code=409, detail="关键词组名称已存在") from exc
        group_id = int(cursor.lastrowid)
        log_event(connection, "keyword.create", f"创建关键词组：{payload.name}")
    return {"id": group_id, "message": "关键词组已创建"}


@app.put("/api/keyword-groups/{group_id}")
def update_keyword_group(group_id: int, payload: KeywordGroupRequest) -> dict[str, Any]:
    validate_id(group_id, "关键词组")
    with get_db() as connection:
        cursor = connection.execute(
            """UPDATE keyword_groups SET name=?,category=?,enabled=?,include_any_json=?,include_all_json=?,phrases_json=?,exclude_json=?,synonyms_json=?,scopes_json=?,priority=?,updated_at=? WHERE id=?""",
            (payload.name, payload.category, int(payload.enabled), json.dumps(payload.include_any, ensure_ascii=False),
             json.dumps(payload.include_all, ensure_ascii=False), json.dumps(payload.phrases, ensure_ascii=False),
             json.dumps(payload.exclude, ensure_ascii=False), json.dumps(payload.synonyms, ensure_ascii=False),
             json.dumps(payload.scopes, ensure_ascii=False), payload.priority, now_iso(), group_id),
        )
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="关键词组不存在")
        rebuilt = rebuild_keyword_group_analysis(connection, group_id)
        log_event(connection, "keyword.update", f"更新关键词组：{payload.name}；已重算 {rebuilt['notice_count']} 条公告，命中 {rebuilt['matched_notice_count']} 条")
    return {"ok": True, "recalculated": rebuilt}


@app.patch("/api/keyword-groups/{group_id}/enabled")
def set_keyword_group_enabled(group_id: int, payload: EnabledRequest) -> dict[str, Any]:
    validate_id(group_id, "关键词组")
    with get_db() as connection:
        row = connection.execute("SELECT name FROM keyword_groups WHERE id=?", (group_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="关键词组不存在")
        connection.execute("UPDATE keyword_groups SET enabled=?,updated_at=? WHERE id=?", (int(payload.enabled), now_iso(), group_id))
        rebuilt = rebuild_keyword_group_analysis(connection, group_id)
        log_event(connection, "keyword.enabled", f"关键词组{('启用' if payload.enabled else '停用')}：{row['name']}；已重算 {rebuilt['notice_count']} 条公告，命中 {rebuilt['matched_notice_count']} 条")
    return {"ok": True, "enabled": payload.enabled, "recalculated": rebuilt}


@app.delete("/api/keyword-groups/{group_id}")
def delete_keyword_group(group_id: int) -> dict[str, Any]:
    validate_id(group_id, "关键词组")
    with get_db() as connection:
        row = connection.execute("SELECT name FROM keyword_groups WHERE id=?", (group_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="关键词组不存在")
        bound_count = connection.execute(
            "SELECT COUNT(*) FROM crawl_jobs WHERE keyword_group_id=?",
            (group_id,),
        ).fetchone()[0]
        if bound_count:
            raise HTTPException(
                status_code=409,
                detail=f"关键词组已绑定 {bound_count} 个采集任务，请先编辑任务解绑后再删除",
            )
        # 命中证据必须与规则绑定；规则删除后不能留下无归属的旧命中。
        connection.execute("DELETE FROM keyword_hits WHERE keyword_group_id=?", (group_id,))
        connection.execute("DELETE FROM keyword_groups WHERE id=?", (group_id,))
        log_event(connection, "keyword.delete", f"删除关键词组：{row['name']}")
    return {"ok": True}


@app.get("/api/crawl-jobs")
def list_jobs() -> dict[str, Any]:
    with get_db() as connection:
        rows = connection.execute("SELECT j.*,s.code AS site_code,s.name AS site_name,g.name AS keyword_group_name FROM crawl_jobs j JOIN sites s ON s.id=j.site_id LEFT JOIN keyword_groups g ON g.id=j.keyword_group_id WHERE s.enabled=1 ORDER BY j.id").fetchall()
    return {"items": [localize({**dict(row), "categories": json_load(row["categories_json"], []), "retry": json_load(row["retry_json"], {})}, ("last_run_at", "schedule_anchor_at", "created_at")) for row in rows]}


@app.post("/api/crawl-jobs")
def create_job(payload: JobRequest) -> dict[str, Any]:
    validate_job_request(payload)
    with get_db() as connection:
        if payload.timezone != "Asia/Shanghai":
            raise HTTPException(status_code=400, detail="系统统一使用北京时间（Asia/Shanghai）")
        if not connection.execute("SELECT id FROM sites WHERE id=?", (payload.site_id,)).fetchone():
            raise HTTPException(status_code=400, detail="目标平台不存在")
        if payload.keyword_group_id and not connection.execute("SELECT id FROM keyword_groups WHERE id=?", (payload.keyword_group_id,)).fetchone():
            raise HTTPException(status_code=400, detail="关键词组不存在")
        timestamp = now_iso()
        cursor = connection.execute("""INSERT INTO crawl_jobs(name,site_id,account_id,keyword_group_id,categories_json,schedule_text,timezone,lookback_days,max_pages,max_notices,concurrency,interval_ms,retry_json,download_attachments,ocr_enabled,enabled,schedule_anchor_at,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (payload.name, payload.site_id, payload.account_id, payload.keyword_group_id, json.dumps(payload.categories, ensure_ascii=False), payload.schedule_text, payload.timezone, payload.lookback_days, payload.max_pages, payload.max_notices, payload.concurrency, payload.interval_ms, json.dumps({"max_attempts": payload.retry_max_attempts}), int(payload.download_attachments), int(payload.ocr_enabled), int(payload.enabled), timestamp if payload.enabled else None, timestamp))
        job_id = int(cursor.lastrowid)
        log_event(connection, "job.create", f"创建采集任务：{payload.name}")
    return {"id": job_id, "message": "采集任务已创建"}


@app.put("/api/crawl-jobs/{job_id}")
def update_job(job_id: int, payload: JobRequest) -> dict[str, Any]:
    validate_id(job_id, "任务")
    validate_job_request(payload)
    with get_db() as connection:
        if payload.timezone != "Asia/Shanghai":
            raise HTTPException(status_code=400, detail="系统统一使用北京时间（Asia/Shanghai）")
        if not connection.execute("SELECT id FROM sites WHERE id=?", (payload.site_id,)).fetchone():
            raise HTTPException(status_code=400, detail="目标平台不存在")
        if payload.keyword_group_id and not connection.execute("SELECT id FROM keyword_groups WHERE id=?", (payload.keyword_group_id,)).fetchone():
            raise HTTPException(status_code=400, detail="关键词组不存在")
        existing = connection.execute("SELECT enabled,schedule_text,schedule_anchor_at FROM crawl_jobs WHERE id=?", (job_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="采集任务不存在")
        schedule_anchor = existing["schedule_anchor_at"]
        if not payload.enabled:
            schedule_anchor = None
        elif not existing["enabled"] or existing["schedule_text"] != payload.schedule_text:
            schedule_anchor = now_iso()
        cursor = connection.execute("""UPDATE crawl_jobs SET name=?,site_id=?,account_id=?,keyword_group_id=?,categories_json=?,schedule_text=?,timezone=?,lookback_days=?,max_pages=?,max_notices=?,concurrency=?,interval_ms=?,retry_json=?,download_attachments=?,ocr_enabled=?,enabled=?,schedule_anchor_at=? WHERE id=?""", (payload.name, payload.site_id, payload.account_id, payload.keyword_group_id, json.dumps(payload.categories, ensure_ascii=False), payload.schedule_text, payload.timezone, payload.lookback_days, payload.max_pages, payload.max_notices, payload.concurrency, payload.interval_ms, json.dumps({"max_attempts": payload.retry_max_attempts}), int(payload.download_attachments), int(payload.ocr_enabled), int(payload.enabled), schedule_anchor, job_id))
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="采集任务不存在")
        log_event(connection, "job.update", f"更新采集任务：{payload.name}")
    return {"ok": True}


@app.patch("/api/crawl-jobs/{job_id}/enabled")
def set_job_enabled(job_id: int, payload: EnabledRequest) -> dict[str, Any]:
    with get_db() as connection:
        row = connection.execute("SELECT j.name,j.enabled,s.enabled AS site_enabled FROM crawl_jobs j JOIN sites s ON s.id=j.site_id WHERE j.id=?", (job_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="采集任务不存在")
        if payload.enabled and not row["site_enabled"]:
            raise HTTPException(status_code=400, detail="该采集平台已停用，不能启用任务")
        anchor = now_iso() if payload.enabled and not row["enabled"] else None if not payload.enabled else connection.execute("SELECT schedule_anchor_at FROM crawl_jobs WHERE id=?", (job_id,)).fetchone()["schedule_anchor_at"]
        connection.execute("UPDATE crawl_jobs SET enabled=?,schedule_anchor_at=? WHERE id=?", (int(payload.enabled), anchor, job_id))
        log_event(connection, "job.enabled", f"任务{('启用' if payload.enabled else '停用')}：{row['name']}")
    return {"ok": True, "enabled": payload.enabled}


@app.delete("/api/crawl-jobs/{job_id}")
def delete_job(job_id: int) -> dict[str, Any]:
    validate_id(job_id, "任务")
    with get_db() as connection:
        row = connection.execute("SELECT name FROM crawl_jobs WHERE id=?", (job_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="采集任务不存在")
        if connection.execute("SELECT 1 FROM crawl_runs WHERE job_id=? AND status IN ('queued','running')", (job_id,)).fetchone():
            raise HTTPException(status_code=409, detail="任务正在运行，不能删除")
        connection.execute("UPDATE system_logs SET crawl_run_id=NULL WHERE crawl_run_id IN (SELECT id FROM crawl_runs WHERE job_id=?)", (job_id,))
        connection.execute("DELETE FROM crawl_jobs WHERE id=?", (job_id,))
        log_event(connection, "job.delete", f"删除采集任务：{row['name']}")
    return {"ok": True}


@app.post("/api/crawl-jobs/run-all")
def run_all_jobs(background_tasks: BackgroundTasks, payload: RunAllJobsRequest | None = None) -> dict[str, Any]:
    with get_db() as connection:
        jobs = connection.execute(
            "SELECT j.id,j.name,j.keyword_group_id,g.enabled AS keyword_enabled "
            "FROM crawl_jobs j LEFT JOIN keyword_groups g ON g.id=j.keyword_group_id "
            "WHERE j.enabled=1 ORDER BY j.id"
        ).fetchall()
    if not jobs:
        raise HTTPException(status_code=409, detail="没有已启用的采集任务")

    overrides = payload.model_dump(exclude_none=True) if payload else {}
    started = []
    skipped = []
    for job in jobs:
        job_id = int(job["id"])
        if job["keyword_group_id"] and not job["keyword_enabled"]:
            skipped.append({"job_id": job_id, "job_name": job["name"], "reason": "绑定的关键词组已停用"})
            continue
        run_id = try_create_run(job_id, reset_schedule=True)
        if run_id is None:
            skipped.append({"job_id": job_id, "job_name": job["name"], "reason": "已有采集批次正在运行"})
            continue
        background_tasks.add_task(run_crawl, job_id, run_id, overrides, True)
        started.append({"job_id": job_id, "job_name": job["name"], "run_id": run_id})

    message = f"已将 {len(started)} 个任务加入采集队列" if started else "没有可启动的新任务"
    if skipped:
        message += f"，跳过 {len(skipped)} 个任务"
    with get_db() as connection:
        log_event(connection, "crawl.run_all", message)
        for item in skipped:
            log_event(connection, "crawl.run_all.skip", f"批量采集跳过 {item['job_name']}：{item['reason']}", "WARNING")
    return {"started": started, "skipped": skipped, "message": message}


@app.post("/api/crawl-jobs/{job_id}/run")
def run_job(job_id: int, background_tasks: BackgroundTasks, payload: RunJobRequest | None = None) -> dict[str, Any]:
    validate_id(job_id, "任务")
    with get_db() as connection:
        job = connection.execute("SELECT j.id,j.name,j.keyword_group_id,g.enabled AS keyword_enabled FROM crawl_jobs j LEFT JOIN keyword_groups g ON g.id=j.keyword_group_id WHERE j.id=? AND j.enabled=1", (job_id,)).fetchone()
        if not job:
            raise HTTPException(status_code=404, detail="任务不存在或已停用")
        if job["keyword_group_id"] and not job["keyword_enabled"]:
            raise HTTPException(status_code=409, detail="任务绑定的关键词组已停用，请先启用或更换关键词组")
        if connection.execute("SELECT 1 FROM crawl_runs WHERE job_id=? AND status IN ('queued','running')", (job_id,)).fetchone():
            raise HTTPException(status_code=409, detail="该任务已有采集批次正在运行")
    # Manual execution starts a fresh countdown immediately at the click time.
    run_id = try_create_run(job_id, reset_schedule=True)
    if run_id is None:
        raise HTTPException(status_code=409, detail="该任务已有采集批次正在运行")
    background_tasks.add_task(run_crawl, job_id, run_id, payload.model_dump(exclude_none=True) if payload else {}, True)
    return {"run_id": run_id, "status": "queued", "message": "采集批次已进入队列，已从本次点击时间重新计时"}


@app.get("/api/crawl-runs/{run_id}")
def get_run(run_id: int) -> dict[str, Any]:
    validate_id(run_id, "采集批次")
    with get_db() as connection:
        row = connection.execute("SELECT r.*,j.name AS job_name FROM crawl_runs r JOIN crawl_jobs j ON j.id=r.job_id WHERE r.id=?", (run_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="采集批次不存在")
    return localize(dict(row), ("started_at", "finished_at", "created_at"))


@app.get("/api/crawl-runs")
def list_runs(limit: int = Query(default=20, ge=1, le=200)) -> dict[str, Any]:
    with get_db() as connection:
        rows = connection.execute("SELECT r.*,j.name AS job_name,s.name AS site_name FROM crawl_runs r JOIN crawl_jobs j ON j.id=r.job_id JOIN sites s ON s.id=j.site_id ORDER BY r.id DESC LIMIT ?", (limit,)).fetchall()
    return {"items": [localize(dict(row), ("started_at", "finished_at", "created_at")) for row in rows]}


@app.get("/api/scheduler/status")
def scheduler_status() -> dict[str, Any]:
    return scheduler.status()


@app.get("/api/dashboard")
def dashboard() -> dict[str, Any]:
    local_now = datetime.now(ZoneInfo("Asia/Shanghai"))
    utc_start = local_now.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat(timespec="seconds")
    utc_end = (local_now.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)).astimezone(timezone.utc).isoformat(timespec="seconds")
    with get_db() as connection:
        counts = connection.execute(f"""SELECT COUNT(*) total, SUM(created_at>=? AND created_at<?) today_new, SUM(business_mark='pending') pending, SUM(business_mark='focus') focus, SUM({PARSE_ISSUE_SQL}) issues, SUM(ingest_status='parsed') parsed FROM notices n WHERE deleted_at IS NULL AND (source_type<>'crawl' OR EXISTS (SELECT 1 FROM keyword_hits h WHERE h.notice_id=n.id AND h.is_negative=0))""", (utc_start, utc_end)).fetchone()
        platforms = [localize(dict(row), ("last_checked_at",)) for row in connection.execute("SELECT s.id,s.code,s.name,s.health_status,s.health_message,s.last_checked_at,COUNT(n.id) notice_count FROM sites s LEFT JOIN notices n ON n.site_id=s.id AND n.deleted_at IS NULL WHERE s.code NOT IN ('ceb','espic') GROUP BY s.id ORDER BY s.id").fetchall()]
        logs = [localize(dict(row), ("created_at",)) for row in connection.execute("SELECT * FROM system_logs ORDER BY id DESC LIMIT 8").fetchall()]
        runs = [localize(dict(row), ("started_at", "finished_at", "created_at")) for row in connection.execute("SELECT r.*,j.name AS job_name FROM crawl_runs r JOIN crawl_jobs j ON j.id=r.job_id ORDER BY r.id DESC LIMIT 8").fetchall()]
    return {"timezone": "Asia/Shanghai", "timezone_label": "北京时间（UTC+08:00）", "counts": {key: int(counts[key] or 0) for key in counts.keys()}, "platforms": platforms, "recent_logs": logs, "recent_runs": runs, "today_definition": "按首次入库时间（Asia/Shanghai）统计"}


@app.get("/api/notices")
def list_notices(q: str = "", platform: str = "all", mark: str = "all", status: str = "all", attachment: str = "all", only_issues: bool = False, only_matched: bool = True, only_unmatched: bool = False, include_deleted: bool = False, only_deleted: bool = False, only_today_new: bool = False, limit: int = Query(default=10, ge=1, le=100), offset: int = Query(default=0, ge=0)) -> dict[str, Any]:
    possible_missed_match = """n.source_type='crawl'
        AND NOT EXISTS (SELECT 1 FROM keyword_hits mh WHERE mh.notice_id=n.id AND mh.is_negative=0)
        AND (n.ingest_status IN ('failed','partial') OR EXISTS (
            SELECT 1 FROM attachments ma WHERE ma.notice_id=n.id
            AND (ma.parse_status IN ('failed','unsupported','ocr_pending','pending')
                 OR ma.status IN ('failed','needs_tool','not_downloaded'))
        ))"""
    clauses = ["1=1"]
    params: list[Any] = []
    local_now = datetime.now(ZoneInfo("Asia/Shanghai"))
    utc_start = local_now.replace(hour=0, minute=0, second=0, microsecond=0).astimezone(timezone.utc).isoformat(timespec="seconds")
    utc_end = (local_now.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1)).astimezone(timezone.utc).isoformat(timespec="seconds")
    if q:
        clauses.append("(n.title LIKE ? OR n.project_number LIKE ? OR n.demand_unit LIKE ? OR n.summary LIKE ? OR EXISTS (SELECT 1 FROM keyword_hits qh WHERE qh.notice_id=n.id AND qh.is_negative=0 AND (qh.keyword LIKE ? OR qh.snippet LIKE ?)))")
        params.extend([f"%{q}%"] * 6)
    if platform != "all":
        clauses.append("s.code=?"); params.append(platform)
    if attachment == "yes":
        clauses.append("EXISTS (SELECT 1 FROM attachments af WHERE af.notice_id=n.id)")
    elif attachment == "no":
        clauses.append("NOT EXISTS (SELECT 1 FROM attachments af WHERE af.notice_id=n.id)")
    base_clauses = list(clauses)
    base_params = list(params)
    if only_deleted:
        clauses.append("n.deleted_at IS NOT NULL")
    elif not include_deleted:
        clauses.append("n.deleted_at IS NULL")
    scope_clauses = base_clauses + ["n.deleted_at IS NULL"]
    scope_params = list(base_params)
    if only_unmatched:
        clauses.append(possible_missed_match)
    elif only_matched:
        clauses.append("(n.source_type<>'crawl' OR EXISTS (SELECT 1 FROM keyword_hits mh WHERE mh.notice_id=n.id AND mh.is_negative=0))")
    # Business-category totals stay on the normal notice-library scope even
    # while the separate possible-missed-match recovery view is selected.
    facet_clauses = scope_clauses + ["(n.source_type<>'crawl' OR EXISTS (SELECT 1 FROM keyword_hits fh WHERE fh.notice_id=n.id AND fh.is_negative=0))"]
    facet_params = scope_params
    if mark != "all":
        clauses.append("n.business_mark=?"); params.append(mark)
    if status != "all":
        clauses.append("n.ingest_status=?"); params.append(status)
    if only_issues:
        clauses.append(PARSE_ISSUE_SQL)
    if only_today_new:
        clauses.append("n.created_at>=? AND n.created_at<?")
        params.extend([utc_start, utc_end])
    order = "n.created_at DESC" if only_today_new else "COALESCE(n.published_at,n.created_at) DESC"
    query = f"SELECT n.* FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE {' AND '.join(clauses)} ORDER BY {order} LIMIT ? OFFSET ?"
    params.extend([limit, offset])
    with get_db() as connection:
        rows = connection.execute(query, params).fetchall()
        total = connection.execute(f"SELECT COUNT(*) FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE {' AND '.join(clauses)}", params[:-2]).fetchone()[0]
        facet_where = " AND ".join(facet_clauses)
        category_counts = connection.execute(
            f"""SELECT COUNT(*) AS all_count,
            SUM(n.business_mark='pending') AS pending_count,
            SUM(n.business_mark='focus') AS focus_count,
            SUM({PARSE_ISSUE_SQL}) AS issues_count,
            SUM(n.created_at>=? AND n.created_at<?) AS today_new_count
            FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE {facet_where}""",
            [utc_start, utc_end] + facet_params,
        ).fetchone()
        unmatched_count = connection.execute(
            f"SELECT COUNT(*) FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE {' AND '.join(scope_clauses)} AND {possible_missed_match}",
            scope_params,
        ).fetchone()[0]
        trash_count = connection.execute(
            f"SELECT COUNT(*) FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE {' AND '.join(base_clauses)} AND n.deleted_at IS NOT NULL",
            base_params,
        ).fetchone()[0]
        items = [frontend_notice(connection, row) for row in rows]
    return {
        "items": items, "total": total, "limit": limit, "offset": offset,
        "category_counts": {
            "all": int(category_counts["all_count"] or 0),
            "pending": int(category_counts["pending_count"] or 0),
            "focus": int(category_counts["focus_count"] or 0),
            "issues": int(category_counts["issues_count"] or 0),
            "unmatched": int(unmatched_count or 0),
            "trash": int(trash_count or 0),
            "today_new": int(category_counts["today_new_count"] or 0),
        },
    }


@app.get("/api/notices/{notice_id}")
def get_notice(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        row = connection.execute("SELECT * FROM notices WHERE id=?", (notice_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="公告不存在")
        item = frontend_notice(connection, row)
        item["versions"] = [localize(dict(version), ("published_at", "captured_at")) for version in connection.execute("SELECT version_no,raw_html_path,published_at,content_fingerprint,captured_at FROM notice_versions WHERE notice_id=? ORDER BY version_no DESC", (notice_id,)).fetchall()]
        item["fields"] = [localize(dict(field), ("created_at", "updated_at")) for field in connection.execute("SELECT * FROM extracted_fields WHERE notice_id=? ORDER BY id", (notice_id,)).fetchall()]
    return item


@app.get("/api/notices/{notice_id}/hits")
def get_notice_hits(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        rows = connection.execute("SELECT h.*,g.name AS keyword_group FROM keyword_hits h JOIN keyword_groups g ON g.id=h.keyword_group_id WHERE h.notice_id=? ORDER BY h.id", (notice_id,)).fetchall()
    return {"items": [localize(dict(row), ("created_at",)) for row in rows]}


def _highlight_evidence(text: str, keyword: str) -> str:
    value = str(text or "")
    needle = str(keyword or "")
    if not value:
        return "<em>暂无可显示的原文内容</em>"
    if not needle:
        return html_escape(value)
    parts: list[str] = []
    cursor = 0
    for index, match in enumerate(re.finditer(re.escape(needle), value, re.IGNORECASE)):
        parts.append(html_escape(value[cursor:match.start()]))
        parts.append(f'<mark id="evidence-hit-{index}">{html_escape(match.group(0))}</mark>')
        cursor = match.end()
        if index >= 20:
            break
    if not parts:
        return html_escape(value)
    parts.append(html_escape(value[cursor:]))
    return "".join(parts)


def _evidence_source(connection, notice_id: int, hit: dict[str, Any]) -> tuple[str, str | None, str | None]:
    source_type = hit["source_type"]
    if source_type == "title":
        notice = connection.execute("SELECT title,source_url FROM notices WHERE id=?", (notice_id,)).fetchone()
        return (notice["title"] if notice else hit["snippet"] or "", notice["source_url"] if notice else None, None)
    if source_type == "body":
        version = connection.execute("SELECT body_text FROM notice_versions WHERE notice_id=? ORDER BY version_no DESC LIMIT 1", (notice_id,)).fetchone()
        notice = connection.execute("SELECT source_url FROM notices WHERE id=?", (notice_id,)).fetchone()
        return (version["body_text"] if version and version["body_text"] else hit["snippet"] or "", notice["source_url"] if notice else None, None)
    file_row = connection.execute(
        "SELECT a.id,a.name,a.relative_path,d.text_content FROM attachments a "
        "LEFT JOIN extracted_documents d ON d.attachment_id=a.id "
        "WHERE a.notice_id=? AND a.name=? ORDER BY d.id DESC LIMIT 1",
        (notice_id, hit["source_file"]),
    ).fetchone()
    if file_row:
        return (file_row["text_content"] or file_row["name"], None, f"/api/files/{file_row['id']}/preview" if file_row["relative_path"] else None)
    return (hit["snippet"] or hit["keyword"], None, None)


@app.get("/api/notices/{notice_id}/evidence/{hit_id}", response_class=HTMLResponse)
def locate_evidence(notice_id: int, hit_id: int) -> HTMLResponse:
    validate_id(notice_id)
    validate_id(hit_id, "命中证据")
    with get_db() as connection:
        notice = connection.execute("SELECT title FROM notices WHERE id=?", (notice_id,)).fetchone()
        hit = connection.execute(
            "SELECT id,keyword,source_type,source_file,location,snippet,context_before,context_after "
            "FROM keyword_hits WHERE id=? AND notice_id=? AND is_negative=0",
            (hit_id, notice_id),
        ).fetchone()
        if not notice or not hit:
            raise HTTPException(status_code=404, detail="命中证据不存在")
        hit = dict(hit)
        content, original_url, file_url = _evidence_source(connection, notice_id, hit)
    context = (hit["context_before"] or "") + (hit["keyword"] or "") + (hit["context_after"] or "")
    display_text = context or content
    source_name = hit["source_file"] or "公告正文.html"
    source_link = f'<a href="{html_escape(file_url, quote=True)}" target="_blank" rel="noreferrer">打开对应文件</a>' if file_url else f'<a href="{html_escape(original_url, quote=True)}" target="_blank" rel="noreferrer">打开原公告</a>' if original_url else ""
    body = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>命中证据定位 - {html_escape(hit['keyword'])}</title>
<style>
body{{margin:0;background:#f5f7fb;color:#24304a;font:14px/1.8 "Microsoft YaHei",sans-serif}}
main{{max-width:960px;margin:32px auto;padding:0 20px}}
.card{{background:#fff;border:1px solid #e4e8f0;border-radius:12px;padding:24px;box-shadow:0 8px 24px rgba(30,45,80,.06)}}
h1{{font-size:20px;margin:0 0 8px}} .meta{{color:#6d7890;font-size:13px;margin-bottom:18px}}
pre{{white-space:pre-wrap;word-break:break-word;background:#f8f9fc;border:1px solid #e9edf4;border-radius:8px;padding:18px;margin:0}}
mark{{background:#ffe58f;color:#5b4300;border-radius:3px;padding:1px 3px;font-weight:700}}
a{{color:#5265d8;text-decoration:none;margin-left:12px}}
</style></head><body><main><div class="card">
<h1>命中证据定位：{html_escape(hit['keyword'])}</h1>
<div class="meta">公告：{html_escape(notice['title'])}<br>来源：{html_escape(source_name)} · {html_escape(hit['location'] or '正文')} {source_link}</div>
<pre id="evidence-hit">{_highlight_evidence(display_text, hit['keyword'])}</pre>
</div></main><script>window.addEventListener("DOMContentLoaded",()=>document.querySelector("mark")?.scrollIntoView({{block:"center"}}));</script></body></html>"""
    return HTMLResponse(body)


@app.get("/api/notices/{notice_id}/files")
def get_notice_files(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        rows = connection.execute("SELECT * FROM attachments WHERE notice_id=? ORDER BY parent_attachment_id IS NOT NULL,id", (notice_id,)).fetchall()
    return {"items": [dict(row) for row in rows]}


@app.patch("/api/files/{file_id}/key")
def mark_key_file(file_id: int, payload: KeyFileRequest) -> dict[str, Any]:
    validate_id(file_id, "文件")
    with get_db() as connection:
        cursor = connection.execute("UPDATE attachments SET is_key_file=? WHERE id=?", (int(payload.is_key_file), file_id))
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="文件不存在")
        row = connection.execute("SELECT notice_id,name FROM attachments WHERE id=?", (file_id,)).fetchone()
        log_event(connection, "attachment.key_file", f"{row['name']}：{'设为' if payload.is_key_file else '取消'}关键文件", notice_id=row["notice_id"])
    return {"ok": True, "is_key_file": payload.is_key_file}


@app.patch("/api/notices/{notice_id}/mark")
def mark_notice(notice_id: int, payload: MarkRequest) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        cursor = connection.execute("UPDATE notices SET business_mark=?,updated_at=? WHERE id=? AND deleted_at IS NULL", (payload.mark, now_iso(), notice_id))
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="公告不存在")
        log_event(connection, "notice.mark", f"公告标记为：{payload.mark}", notice_id=notice_id)
    return {"ok": True, "mark": payload.mark}


@app.patch("/api/notices/{notice_id}/fields/{field_name}")
def correct_notice_field(notice_id: int, field_name: str, payload: FieldCorrectionRequest) -> dict[str, Any]:
    allowed = {"project_number", "project_name", "demand_unit", "procuring_agent", "project_location", "budget"}
    if field_name not in allowed:
        raise HTTPException(status_code=400, detail="该字段不允许人工修订")
    with get_db() as connection:
        if not connection.execute("SELECT id FROM notices WHERE id=?", (notice_id,)).fetchone():
            raise HTTPException(status_code=404, detail="公告不存在")
        connection.execute(f"UPDATE notices SET {field_name}=?,updated_at=? WHERE id=?", (payload.value, now_iso(), notice_id))
        connection.execute(
            "INSERT INTO extracted_fields(notice_id,field_name,value,source_location,extraction_method,confidence,manual_value,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
            (notice_id, field_name, payload.value, "人工修订", "manual", 1.0, payload.value, now_iso(), now_iso()),
        )
        log_event(connection, "notice.field_correct", f"人工修订字段：{field_name}", notice_id=notice_id)
    return {"ok": True, "field": field_name, "value": payload.value}


@app.post("/api/notices/batch-mark")
def batch_mark(payload: BatchActionRequest) -> dict[str, Any]:
    if not payload.mark:
        raise HTTPException(status_code=400, detail="缺少标记值")
    placeholders = ",".join("?" for _ in payload.ids)
    with get_db() as connection:
        cursor = connection.execute(f"UPDATE notices SET business_mark=?,updated_at=? WHERE id IN ({placeholders}) AND deleted_at IS NULL", [payload.mark, now_iso(), *payload.ids])
        log_event(connection, "notice.batch_mark", f"批量标记 {cursor.rowcount} 条公告为 {payload.mark}")
    return {"ok": True, "updated": cursor.rowcount}


@app.post("/api/notices/batch-delete")
def batch_delete(payload: BatchActionRequest) -> dict[str, Any]:
    placeholders = ",".join("?" for _ in payload.ids)
    timestamp = now_iso()
    with get_db() as connection:
        cursor = connection.execute(f"UPDATE notices SET deleted_at=?,updated_at=? WHERE id IN ({placeholders}) AND deleted_at IS NULL", [timestamp, timestamp, *payload.ids])
        log_event(connection, "notice.batch_delete", f"批量移入回收站 {cursor.rowcount} 条公告")
    return {"ok": True, "deleted": cursor.rowcount}


@app.post("/api/notices/batch-permanent-delete")
def batch_permanent_delete(payload: BatchActionRequest) -> dict[str, Any]:
    if not payload.ids:
        return {"ok": True, "deleted": 0}
    with get_db() as connection:
        deleted = permanently_delete_notices(connection, payload.ids)
        log_event(connection, "notice.permanent_delete", f"永久删除回收站公告 {deleted} 条")
    return {"ok": True, "deleted": deleted}


@app.post("/api/notices/trash/empty")
def empty_trash() -> dict[str, Any]:
    with get_db() as connection:
        ids = [row["id"] for row in connection.execute("SELECT id FROM notices WHERE deleted_at IS NOT NULL").fetchall()]
        deleted = permanently_delete_notices(connection, ids, reason="empty_trash")
        log_event(connection, "notice.trash_empty", f"清空回收站，永久删除 {deleted} 条公告")
    return {"ok": True, "deleted": deleted}


@app.get("/api/notices-export.csv")
def export_notices() -> StreamingResponse:
    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(["平台", "公告标题", "公告类型", "发布时间", "开标时间", "项目编号", "需求单位", "业务标记", "原网页链接"])
    with get_db() as connection:
        rows = connection.execute("SELECT s.name,n.* FROM notices n LEFT JOIN sites s ON s.id=n.site_id WHERE n.deleted_at IS NULL ORDER BY n.id DESC").fetchall()
        for row in rows:
            writer.writerow([row["name"] or "人工导入", row["title"], row["notice_type"], row["published_at"], row["opening_at"], row["project_number"], row["demand_unit"], row["business_mark"], row["source_url"]])
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=notices.csv"})


@app.delete("/api/notices/{notice_id}")
def delete_notice(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        cursor = connection.execute("UPDATE notices SET deleted_at=?,updated_at=? WHERE id=? AND deleted_at IS NULL", (now_iso(), now_iso(), notice_id))
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="公告不存在或已经删除")
        log_event(connection, "notice.soft_delete", "公告已移入回收站", notice_id=notice_id)
    return {"ok": True, "message": "公告已移入回收站"}


@app.delete("/api/notices/{notice_id}/permanent")
def permanently_delete_notice(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        deleted = permanently_delete_notices(connection, [notice_id])
        if not deleted:
            raise HTTPException(status_code=404, detail="公告不存在或不在回收站")
        log_event(connection, "notice.permanent_delete", "公告已永久删除")
    return {"ok": True, "message": "公告及其附件已永久删除"}


@app.post("/api/notices/{notice_id}/restore")
def restore_notice(notice_id: int) -> dict[str, Any]:
    validate_id(notice_id)
    with get_db() as connection:
        cursor = connection.execute("UPDATE notices SET deleted_at=NULL,updated_at=? WHERE id=?", (now_iso(), notice_id))
        if not cursor.rowcount:
            raise HTTPException(status_code=404, detail="公告不存在")
    return {"ok": True}


@app.post("/api/notices/{notice_id}/reparse")
def reparse(notice_id: int, background_tasks: BackgroundTasks, ocr: bool = False) -> dict[str, Any]:
    validate_id(notice_id)
    try:
        if ocr:
            from .ocr import available
            if not available():
                raise HTTPException(status_code=503, detail="未安装 CPU OCR 组件，请使用 Python 3.9-3.13 安装 PaddlePaddle 与 PaddleOCR")
        task_id, created = reparse_tasks.queue(notice_id, use_ocr=ocr)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    if created:
        reparse_tasks.submit(task_id)
    label = "OCR 解析" if ocr else "重新解析"
    return {"ok": True, "task_id": task_id, "status": "queued", "message": f"公告已进入{label}队列" if created else "该公告正在解析，无需重复提交"}


@app.get("/api/reparse-tasks/{task_id}")
def reparse_status(task_id: int):
    task = reparse_tasks.get(task_id)
    if task is None:
        raise HTTPException(status_code=404, detail="解析任务不存在")
    return task


@app.get("/api/notices/{notice_id}/reparse-status")
def latest_reparse_status(notice_id: int):
    reparse_tasks.initialize()
    with get_db() as c:
        task = c.execute('SELECT * FROM reparse_tasks WHERE notice_id=? ORDER BY id DESC LIMIT 1', (notice_id,)).fetchone()
        return {'task': dict(task) if task else None}


@app.post("/api/imports/url")
def import_urls(payload: UrlImportRequest) -> dict[str, Any]:
    results = {"created": 0, "updated": 0, "failed": 0, "items": []}
    with get_db() as connection:
        cursor = connection.execute("INSERT INTO import_batches(source_type,source_name,total_count,status,created_at) VALUES(?,?,?,?,?)", ("url", "批量 URL 导入", len(payload.urls), "running", now_iso()))
        batch_id = int(cursor.lastrowid)
    for source_url in payload.urls:
        source_url = source_url.strip()
        if not source_url.startswith(("http://", "https://")):
            results["failed"] += 1; results["items"].append({"url": source_url, "status": "格式错误", "error": "仅支持 http/https URL"}); continue
        try:
            notice_id, created = ingest_url(source_url, payload.site_code)
            key = "created" if created else "updated"
            results[key] += 1; results["items"].append({"url": source_url, "status": "新增入库" if created else "更新已有公告", "notice_id": notice_id})
        except Exception as exc:
            results["failed"] += 1; results["items"].append({"url": source_url, "status": "失败", "error": str(exc)})
    with get_db() as connection:
        connection.execute("UPDATE import_batches SET status=?,created_count=?,updated_count=?,error_count=? WHERE id=?", ("completed" if not results["failed"] else "partial", results["created"], results["updated"], results["failed"], batch_id))
        log_event(connection, "import.url", f"URL 导入完成：新增 {results['created']} 条，更新 {results['updated']} 条，失败 {results['failed']} 条")
    return {"batch_id": batch_id, **results}


@app.post("/api/imports/excel")
async def import_excel(file: UploadFile = File(...), site_code: str = Form("")) -> dict[str, Any]:
    if not file.filename or Path(file.filename).suffix.lower() not in {".xlsx", ".xls"}:
        raise HTTPException(status_code=400, detail="请上传 xlsx 或 xls 文件")
    if Path(file.filename).suffix.lower() == ".xls":
        raise HTTPException(status_code=400, detail="V1 请先将旧版 xls 另存为 xlsx")
    data = await file.read()
    temp_path = settings.temp_dir / safe_name(file.filename)
    temp_path.write_bytes(data)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(temp_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(status_code=400, detail="Excel 没有可导入数据")
        headers = [str(value or "").strip() for value in rows[0]]
        index = {header: position for position, header in enumerate(headers)}
        title_key = next((key for key in ("公告标题", "标题", "项目名称") if key in index), None)
        if not title_key:
            raise HTTPException(status_code=400, detail="缺少公告标题列")
        created = 0
        for row_number, row in enumerate(rows[1:], start=2):
            title = str(row[index[title_key]] or "").strip()
            if not title:
                continue
            def value(*names: str) -> str | None:
                for name in names:
                    if name in index and index[name] < len(row) and row[index[name]] is not None:
                        return str(row[index[name]]).strip()
                return None
            source_url = value("原网页链接", "来源链接", "URL") or f"manual://excel/{file.filename}/{row_number}"
            notice = NoticeData(external_id=f"excel-{file.filename}-{row_number}", title=title, url=source_url, body_text="；".join(filter(None, [value("项目摘要", "摘要"), value("需求单位", "采购人"), value("项目编号")])), published_at=value("发布时间"), opening_at=value("开标时间", "投标截止时间"), notice_type=value("公告类型") or "人工导入", raw_html="")
            with get_db() as connection:
                site = connection.execute("SELECT id FROM sites WHERE code=?", (site_code,)).fetchone() if site_code else None
                ingest_notice_data(connection, site["id"] if site else None, notice, source_type="excel_import", download_attachments=False)
            created += 1
        workbook.close()
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return {"status": "completed", "created": created, "message": f"Excel 导入完成，共新增 {created} 条"}


@app.post("/api/imports/file")
async def import_file(file: UploadFile = File(...), source_url: str = Form(""), site_code: str = Form("")) -> dict[str, Any]:
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名不能为空")
    content = await file.read()
    if len(content) > settings.max_attachment_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"文件超过 {settings.max_attachment_mb} MB 限制")
    title = Path(file.filename).stem
    url = source_url.strip() or f"manual://file/{file.filename}"
    body = ""
    temp_path = settings.temp_dir / safe_name(file.filename)
    temp_path.write_bytes(content)
    result = parse_document(temp_path)
    body = result.text
    with get_db() as connection:
        site = connection.execute("SELECT id,code,base_url FROM sites WHERE code=?", (site_code,)).fetchone() if site_code else None
        notice_id = ingest_notice_data(connection, site["id"] if site else None, NoticeData(external_id=f"file-{file.filename}-{len(content)}", title=title, url=url, body_text=body, raw_html=body, notice_type="文件导入"), source_type="file_import", download_attachments=False)
        attachment_id = create_attachment(connection, notice_id, file.filename, None, "stored")
        target = attachment_directory(notice_id) / safe_name(file.filename)
        target.write_bytes(content)
        connection.execute("UPDATE attachments SET relative_path=?,mime_type=?,sha256=?,size_bytes=?,parse_status=?,error_message=? WHERE id=?", (relative_to_data(target), file_mime(target), sha256_file(target), len(content), result.status, result.error, attachment_id))
        connection.execute("INSERT INTO extracted_documents(attachment_id,text_content,structure_json,parser,status,created_at,updated_at) VALUES(?,?,?,?,?,?,?)", (attachment_id, result.text, json.dumps(result.structure, ensure_ascii=False), result.parser, result.status, now_iso(), now_iso()))
        refresh_notice_analysis(connection, notice_id, title, body)
        log_event(connection, "import.file", f"文件导入：{file.filename}", notice_id=notice_id)
    if temp_path.exists():
        temp_path.unlink()
    return {"status": "completed", "notice_id": notice_id, "file_id": attachment_id}


@app.get("/api/files/{file_id}/preview")
def preview_file(file_id: int) -> FileResponse:
    validate_id(file_id, "文件")
    with get_db() as connection:
        row = connection.execute("SELECT name,relative_path,mime_type FROM attachments WHERE id=?", (file_id,)).fetchone()
    if not row or not row["relative_path"]:
        raise HTTPException(status_code=404, detail="文件尚未落盘")
    try:
        path = absolute_from_relative(row["relative_path"])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not path.is_file():
        raise HTTPException(status_code=404, detail="文件不存在")
    return FileResponse(path, media_type=row["mime_type"] or "application/octet-stream", filename=row["name"], content_disposition_type="inline")


def _launch_file_location(path: Path) -> str:
    """Open the local file in the OS file manager and select it when supported."""
    if sys.platform.startswith("win"):
        subprocess.Popen(["explorer.exe", f"/select,{path}"])
        return "select"
    if sys.platform == "darwin":
        subprocess.Popen(["open", "-R", str(path)])
        return "select"
    # When the development server runs inside WSL, prefer Windows Explorer so
    # the user sees the same folder as the Windows desktop application.
    explorer = shutil.which("explorer.exe")
    if explorer:
        windows_path = str(path)
        wslpath = shutil.which("wslpath")
        if wslpath:
            try:
                windows_path = subprocess.check_output(
                    [wslpath, "-w", str(path)], text=True, stderr=subprocess.DEVNULL
                ).strip() or windows_path
            except (OSError, subprocess.SubprocessError):
                pass
        subprocess.Popen([explorer, f"/select,{windows_path}"])
        return "select"
    opener = shutil.which("xdg-open")
    if opener:
        subprocess.Popen([opener, str(path.parent)])
        return "directory"
    raise OSError("当前系统没有可用的文件管理器")


@app.api_route("/api/files/{file_id}/open-location", methods=["GET", "POST"])
def open_file_location(file_id: int) -> dict[str, Any]:
    validate_id(file_id, "文件")
    with get_db() as connection:
        row = connection.execute("SELECT name,relative_path FROM attachments WHERE id=?", (file_id,)).fetchone()
    if not row or not row["relative_path"]:
        raise HTTPException(status_code=404, detail="文件尚未落盘")
    try:
        path = absolute_from_relative(row["relative_path"])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not path.is_file():
        raise HTTPException(status_code=404, detail="文件不存在")
    try:
        mode = _launch_file_location(path)
    except (OSError, subprocess.SubprocessError) as exc:
        raise HTTPException(status_code=503, detail=f"无法打开本地文件夹：{exc}") from exc
    message = "已在文件夹中定位文件" if mode == "select" else "已打开文件所在目录"
    return {"ok": True, "mode": mode, "name": row["name"], "path": str(path), "directory": str(path.parent), "message": message}


@app.get("/api/logs")
def list_logs(limit: int = Query(default=100, ge=1, le=500)) -> dict[str, Any]:
    with get_db() as connection:
        rows = connection.execute("SELECT * FROM system_logs ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    return {"items": [localize(dict(row), ("created_at",)) for row in rows]}


# API 路由注册完成后再挂载前端，单机部署无需额外 Nginx。
@app.get("/", include_in_schema=False)
def frontend_index():
    return FileResponse(BASE_DIR / "index.html")


@app.get("/{asset_name}", include_in_schema=False)
def frontend_asset(asset_name: str):
    if asset_name not in {"app.js", "styles.css", "index.html"}:
        raise HTTPException(status_code=404, detail="文件不存在")
    return FileResponse(BASE_DIR / asset_name)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("backend.main:app", host=settings.host, port=settings.port, reload=False)
