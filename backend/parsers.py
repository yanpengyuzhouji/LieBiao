from __future__ import annotations

import base64
import hashlib
import html
import io
import mimetypes
import os
import subprocess
import tempfile
import quopri
import re
import shutil
import zipfile
import zlib
import threading
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from typing import Any


class TextCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        cleaned = re.sub(r"\s+", " ", data).strip()
        if cleaned:
            self.parts.append(cleaned)


@dataclass
class DocumentResult:
    text: str = ""
    structure: dict[str, Any] = field(default_factory=dict)
    parser: str = "none"
    status: str = "unsupported"
    error: str | None = None


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def file_mime(path: Path) -> str:
    guessed, _ = mimetypes.guess_type(path.name)
    if guessed:
        return guessed
    suffix = path.suffix.lower()
    return {
        ".md": "text/markdown", ".html": "text/html", ".htm": "text/html", ".xml": "application/xml",
        ".doc": "application/msword", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xls": "application/vnd.ms-excel", ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".zip": "application/zip", ".rar": "application/vnd.rar", ".7z": "application/x-7z-compressed",
    }.get(suffix, "application/octet-stream")


def decode_text(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_text_like(path: Path) -> DocumentResult:
    raw = path.read_bytes()
    text = decode_text(raw)
    if path.suffix.lower() in {".html", ".htm", ".xml"}:
        collector = TextCollector()
        try:
            collector.feed(text)
            text = "\n".join(collector.parts)
        except Exception:
            text = re.sub(r"<[^>]+>", " ", text)
        return DocumentResult(text=text, structure={"format": path.suffix.lower()}, parser="html", status="parsed")
    return DocumentResult(text=text, structure={"format": path.suffix.lower()}, parser="text", status="parsed")


def parse_pdf(path: Path) -> DocumentResult:
    try:
        import fitz  # PyMuPDF

        pages: list[dict[str, Any]] = []
        document = fitz.open(path)
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text("text")
            pages.append({"page": page_number, "text": page_text})
        document.close()
        return DocumentResult(
            text="\n".join(item["text"] for item in pages),
            structure={"pages": pages},
            parser="pymupdf",
            status="parsed" if any(item["text"].strip() for item in pages) else "ocr_pending",
            error=None if any(item["text"].strip() for item in pages) else "扫描 PDF 未提取到文本，需要启用 OCR",
        )
    except Exception as exc:
        return DocumentResult(parser="pymupdf", status="failed", error=f"PDF 解析失败：{exc}")


def parse_docx(path: Path) -> DocumentResult:
    try:
        from docx import Document

        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        tables: list[dict[str, Any]] = []
        for table_index, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append({"table": table_index, "rows": rows})
        table_text = [" | ".join(row) for table in tables for row in table["rows"]]
        return DocumentResult(text="\n".join(paragraphs + table_text), structure={"paragraphs": paragraphs, "tables": tables}, parser="python-docx", status="parsed")
    except Exception as exc:
        return DocumentResult(parser="python-docx", status="failed", error=f"Word 解析失败：{exc}")


def parse_mime_html_workbook(path: Path) -> DocumentResult:
    raw = path.read_bytes()
    decoded = decode_text(quopri.decodestring(raw))
    html_start = decoded.lower().find("<html")
    if html_start < 0:
        return DocumentResult(parser="mime-html", status="failed", error="MIME 工作表中未找到 HTML 内容")
    collector = TextCollector()
    try:
        collector.feed(decoded[html_start:])
    except Exception as exc:
        return DocumentResult(parser="mime-html", status="failed", error=f"MIME 工作表解析失败：{exc}")
    text = "\n".join(collector.parts)
    return DocumentResult(
        text=text, structure={"format": "mime-html-workbook"}, parser="mime-html",
        status="parsed" if text.strip() else "failed",
        error=None if text.strip() else "MIME 工作表未提取到文本",
    )


def parse_xlsx(path: Path) -> DocumentResult:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        all_text: list[str] = []
        for sheet in workbook.worksheets:
            rows: list[dict[str, Any]] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append({"row": row_number, "values": values})
                    all_text.append(" | ".join(values))
            sheets.append({"name": sheet.title, "rows": rows})
        workbook.close()
        return DocumentResult(text="\n".join(all_text), structure={"sheets": sheets}, parser="openpyxl", status="parsed")
    except Exception as exc:
        header = path.read_bytes()[:512].lstrip()
        if header.startswith(b"MIME-Version:") or b"Content-Type: text/html" in header:
            return parse_mime_html_workbook(path)
        return DocumentResult(parser="openpyxl", status="failed", error=f"Excel 解析失败：{exc}")


def _windows_path(path: Path) -> str:
    value = str(path.resolve())
    match = re.match(r"^/mnt/([a-zA-Z])/(.*)$", value)
    if match:
        return f"{match.group(1).upper()}:\\{match.group(2).replace(chr(47), chr(92))}"
    return value


_word_lock = threading.Lock()


def parser_capabilities():
    import importlib.util
    word = '未检测；旧版 DOC 需要本机 Microsoft Word'
    if os.name == 'nt':
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r'Word.Application\CLSID'):
                word = '检测到 Word 注册信息；实际可用性以文件解析为准'
        except OSError:
            word = '未检测到 Microsoft Word，旧版 DOC 暂不可解析'
    available = lambda module: importlib.util.find_spec(module) is not None
    return {
        'doc': word,
        'docx': '可用' if available('docx') else '缺少解析组件',
        'xlsx': '可用' if available('openpyxl') else '缺少解析组件',
        'pdf': '文本 PDF 可用' if available('fitz') else '缺少解析组件',
        'ocr': '尚未接入，扫描 PDF 和图片需人工处理',
    }


def parse_legacy_doc(path: Path) -> DocumentResult:
    with _word_lock:
        return _parse_legacy_doc(path)


def _parse_legacy_doc(path: Path) -> DocumentResult:
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    bundled = Path("/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe")
    if not powershell and bundled.exists():
        powershell = str(bundled)
    if not powershell:
        return DocumentResult(parser="word-com", status="unsupported", error="未找到 PowerShell，无法调用本机 Word 解析旧版 DOC")

    handle = tempfile.NamedTemporaryFile(prefix="liebiao-doc-", suffix=".txt", dir=path.parent, delete=False)
    output_path = Path(handle.name)
    handle.close()
    input_literal = _windows_path(path).replace("'", "''")
    output_literal = _windows_path(output_path).replace("'", "''")
    pid_path = output_path.with_suffix('.pid')
    pid_literal = _windows_path(pid_path).replace("'", "''")
    command = (
        "$ErrorActionPreference='Stop';"
        "$prior=@(Get-Process WINWORD -ErrorAction SilentlyContinue | Select-Object -ExpandProperty Id);"
        "Add-Type -TypeDefinition 'using System; using System.Runtime.InteropServices; public class WordPid { [DllImport(\"user32.dll\")] public static extern uint GetWindowThreadProcessId(IntPtr h, out uint p); }';"
        f"$inputPath='{input_literal}';$outputPath='{output_literal}';$word=$null;$doc=$null;$owned=$false;try{{"
        "$word=New-Object -ComObject Word.Application;"
        "$word.Visible=$false;$word.DisplayAlerts=0;$word.AutomationSecurity=3;"
        "$probe=$word.Documents.Add();"
        "try{[uint32]$wordPid=0;[void][WordPid]::GetWindowThreadProcessId([IntPtr]$probe.ActiveWindow.Hwnd,[ref]$wordPid)}finally{$probe.Close($false)};"
        "if(!$wordPid -or $prior -contains $wordPid){throw '无法隔离 Word 实例，请稍后重试'};"
        "$owned=$true;$process=Get-Process -Id $wordPid;"
        f"[IO.File]::WriteAllText('{pid_literal}',($wordPid.ToString()+','+$process.StartTime.Ticks));"
        "$word.Visible=$false;$word.DisplayAlerts=0;$word.AutomationSecurity=3;"
        "$doc=$word.Documents.Open($inputPath,$false,$true);"
        "[IO.File]::WriteAllText($outputPath,$doc.Content.Text,[Text.Encoding]::UTF8)"
        "}finally{if($doc){$doc.Close($false)};if($word -and $owned){$word.Quit()}}"
    )
    environment = os.environ.copy()
    environment["LIEBIAO_DOC_INPUT"] = _windows_path(path)
    environment["LIEBIAO_DOC_OUTPUT"] = _windows_path(output_path)
    if os.name != "nt":
        inherited = environment.get("WSLENV", "")
        forwarded = "LIEBIAO_DOC_INPUT/u:LIEBIAO_DOC_OUTPUT/u"
        environment["WSLENV"] = f"{inherited}:{forwarded}" if inherited else forwarded
    encoded_command = base64.b64encode(command.encode("utf-16le")).decode("ascii")
    try:
        completed = subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded_command],
            env=environment, capture_output=True, timeout=45, check=False,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0,
        )
        if completed.returncode != 0:
            detail = decode_text(completed.stderr or completed.stdout).strip()
            return DocumentResult(parser="word-com", status="failed", error=f"旧版 Word 解析失败：{detail or completed.returncode}")
        text = output_path.read_text(encoding="utf-8-sig", errors="replace")
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
        text = re.sub(r"[ \t]+", " ", text).strip()
        return DocumentResult(
            text=text, structure={"format": "word-97-2003"}, parser="word-com",
            status="parsed" if text else "failed",
            error=None if text else "旧版 Word 文档未提取到文本",
        )
    except subprocess.TimeoutExpired:
        return DocumentResult(parser="word-com", status="failed", error="旧版 Word 解析超时（45 秒）")
    except Exception as exc:
        return DocumentResult(parser="word-com", status="failed", error=f"旧版 Word 解析失败：{exc}")
    finally:
        if pid_path.exists():
            try:
                process_id, start_ticks = pid_path.read_text(encoding='utf-8-sig').strip().split(',')
                if process_id.isdigit() and start_ticks.isdigit():
                    cleanup = f"$p=Get-Process -Id {process_id} -ErrorAction SilentlyContinue;if($p -and $p.ProcessName -eq 'WINWORD' -and $p.StartTime.Ticks -eq {start_ticks}){{Stop-Process -Id {process_id} -Force}}"
                    subprocess.run([powershell, '-NoProfile', '-NonInteractive', '-EncodedCommand', base64.b64encode(cleanup.encode('utf-16le')).decode('ascii')], capture_output=True, timeout=10, creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            except (OSError, ValueError, subprocess.SubprocessError):
                pass
            finally:
                pid_path.unlink(missing_ok=True)
        try:
            output_path.unlink(missing_ok=True)
        except OSError:
            pass


def parse_document(path: Path) -> DocumentResult:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv", ".log", ".md", ".html", ".htm", ".xml"}:
        return parse_text_like(path)
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".xlsx":
        return parse_xlsx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        return DocumentResult(parser="ocr", status="ocr_pending", error="图片文件待接入 OCR 引擎")
    if suffix == ".doc":
        return parse_legacy_doc(path)
    if suffix == ".xls":
        return DocumentResult(parser="office-converter", status="unsupported", error="旧版 Excel 文件需安装转换组件")
    return DocumentResult(parser="none", status="unsupported", error=f"暂不支持 {suffix or '无扩展名'} 文件")


def _safe_member_path(destination: Path, member_name: str) -> Path:
    normalized = member_name.replace("\\", "/")
    pure = PurePosixPath(normalized)
    if pure.is_absolute() or ".." in pure.parts:
        raise ValueError(f"压缩包包含不安全路径：{member_name}")
    target = (destination / Path(*pure.parts)).resolve()
    destination_resolved = destination.resolve()
    try:
        target.relative_to(destination_resolved)
    except ValueError as exc:
        raise ValueError(f"压缩包路径越界：{member_name}") from exc
    return target


def is_office_lock_file(name: str) -> bool:
    return name.replace("\\", "/").rsplit("/", 1)[-1].startswith("~$")


def safe_extract_zip(path: Path, destination: Path, max_archive_bytes: int, max_expanded_bytes: int, max_files: int, max_depth: int) -> list[Path]:
    if path.stat().st_size > max_archive_bytes:
        raise ValueError(f"压缩包超过大小限制（{max_archive_bytes // (1024 * 1024)} MB）")
    destination.mkdir(parents=True, exist_ok=True)
    extracted: list[Path] = []
    expanded_bytes = 0
    try:
        archive = zipfile.ZipFile(path)
        members_probe = archive.infolist()
        repair_needed = False
        for member in members_probe:
            if member.flag_bits & 0x800:
                continue
            try:
                decoded = member.filename.encode("cp437").decode("utf-8")
                repair_needed = repair_needed or decoded != member.filename
            except (UnicodeEncodeError, UnicodeDecodeError):
                pass
        if repair_needed:
            archive.close()
            raise zipfile.BadZipFile("filename encoding flag differ")
    except zipfile.BadZipFile as exc:
        # 部分公告 ZIP 的文件名实际为 UTF-8，但生成器遗漏了 UTF-8 标志。
        # 只修正内存中的解析副本，原始附件保持不变。
        if "differ" not in str(exc):
            raise
        repaired = bytearray(path.read_bytes())
        for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            start = 0
            while True:
                position = repaired.find(signature, start)
                if position < 0:
                    break
                offset = position + flag_offset
                flags = int.from_bytes(repaired[offset:offset + 2], "little") | 0x800
                repaired[offset:offset + 2] = flags.to_bytes(2, "little")
                start = position + 4
        archive = zipfile.ZipFile(io.BytesIO(repaired))
    with archive:
        members = archive.infolist()
        for member in members:
            if member.flag_bits & 0x800:
                continue
            central_name = member.filename
            try:
                corrected_name = central_name.encode("cp437").decode("gb18030")
            except (UnicodeEncodeError, UnicodeDecodeError):
                corrected_name = central_name
            # 某些国网 ZIP 的中央目录用 GBK、本地文件头用 UTF-8，且均未设编码位。
            # 记录本地头的实际名称供 zipfile 一致性校验，展示/落盘则用修正后名称。
            if archive.fp is not None:
                position = archive.fp.tell()
                archive.fp.seek(member.header_offset + 6)
                local_flags_raw = archive.fp.read(2)
                archive.fp.seek(member.header_offset + 26)
                lengths = archive.fp.read(4)
                if len(lengths) == 4:
                    name_length = int.from_bytes(lengths[:2], "little")
                    extra_length = int.from_bytes(lengths[2:], "little")
                    archive.fp.seek(member.header_offset + 30)
                    local_name_raw = archive.fp.read(name_length)
                    try:
                        local_flags = int.from_bytes(local_flags_raw, "little") if len(local_flags_raw) == 2 else 0
                        member.orig_filename = local_name_raw.decode("utf-8" if local_flags & 0x800 else "cp437")
                    except UnicodeDecodeError:
                        pass
                archive.fp.seek(position)
            member.filename = corrected_name
        if len(members) > max_files:
            raise ValueError(f"压缩包文件数量超过限制（{max_files} 个）")
        for member in members:
            if member.is_dir():
                continue
            if is_office_lock_file(member.filename):
                continue
            depth = len([part for part in member.filename.replace("\\", "/").split("/") if part])
            if depth > max_depth + 1:
                raise ValueError(f"压缩包目录层数超过限制（{max_depth} 层）")
            expanded_bytes += member.file_size
            if expanded_bytes > max_expanded_bytes:
                raise ValueError(f"压缩包展开后超过限制（{max_expanded_bytes // (1024 * 1024)} MB）")
            if member.compress_size and member.file_size / member.compress_size > 200:
                raise ValueError(f"压缩包疑似高压缩比文件：{member.filename}")
            target = _safe_member_path(destination, member.filename)
            target.parent.mkdir(parents=True, exist_ok=True)
            # Reparse must be idempotent.  Reuse an already extracted member
            # when its size is unchanged; Windows/Word may temporarily hold a
            # legacy DOC open and reject an otherwise unnecessary overwrite.
            if target.is_file() and target.stat().st_size == member.file_size:
                checksum = 0
                with target.open("rb") as existing:
                    for chunk in iter(lambda: existing.read(1024 * 1024), b""):
                        checksum = zlib.crc32(chunk, checksum)
                if checksum == member.CRC:
                    extracted.append(target)
                    continue
            temporary = None
            try:
                with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".extract-", delete=False) as sink:
                    temporary = Path(sink.name)
                    with archive.open(member) as source:
                        shutil.copyfileobj(source, sink, length=1024 * 1024)
                temporary.replace(target)
            finally:
                if temporary is not None:
                    temporary.unlink(missing_ok=True)
            extracted.append(target)
    return extracted


def archive_type(path: Path) -> bool:
    return path.suffix.lower() in {".zip", ".rar", ".7z"}
